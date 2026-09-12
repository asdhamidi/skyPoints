"""
Converts AUS.xlsx to CSV ahead of staging — Snowflake has no native Excel file
format (docs/02_Tech_Stack_Specification.md §2). Runs as the convert_aus_xlsx
Airflow task, before stage_and_copy_aus (docs/03_Technical_Build_Plan.md §4).

Text cells — including the literal "NULL" string and any invalid date string
(docs/01_Design_Specification.md §2, §4) — pass through unchanged: NULL_IF at
the Snowflake file-format level (setup/snowflake_bootstrap.sql §4) is what
turns "NULL" into a true null, not this script. Real datetime cells are
written out in a fixed ISO 8601 format (YYYY-MM-DD) so downstream parsing has
exactly one format to handle for genuine dates.

Public API:
    convert(xlsx_path: Path, csv_path: Path) -> Path

CLI:
    python convert_aus_xlsx_to_csv.py landing/aus_member_20240115.xlsx landing/aus_member_20240115.csv
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
from pathlib import Path

from openpyxl import load_workbook


def _format_cell(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    if value is None:
        return ""
    return value


def convert(xlsx_path: Path, csv_path: Path) -> Path:
    xlsx_path = Path(xlsx_path)
    csv_path = Path(csv_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(f"source workbook not found: {xlsx_path}")

    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow([_format_cell(v) for v in row])

    return csv_path


def _parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx_path", type=Path)
    parser.add_argument("csv_path", type=Path)
    return parser.parse_args()


def main() -> None:
    args = _parse_args()
    result = convert(args.xlsx_path, args.csv_path)
    print(result)


if __name__ == "__main__":
    main()
